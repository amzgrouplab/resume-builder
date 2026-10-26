from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.parse_resume import parse_document
from app.write_resume import apply_bullet_updates


def _sample_doc(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    name = doc.add_paragraph()
    run = name.add_run("Alex Rivera")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading = doc.add_paragraph()
    hrun = heading.add_run("EXPERIENCE")
    hrun.bold = True
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(12)

    role = doc.add_paragraph()
    rrun = role.add_run("Staff Engineer, Northwind  |  Jan 2018 – Dec 2020")
    rrun.bold = True
    rrun.font.name = "Calibri"
    rrun.font.size = Pt(11)

    bullets = [
        "Built Java 8 Spring Boot services for claims intake and reduced batch failures across nightly jobs.",
        "Designed PostgreSQL schemas and indexed queries that cut report time for finance operations.",
        "Automated Jenkins pipelines for test, package, and deploy of internal services.",
    ]
    for text in bullets:
        paragraph = doc.add_paragraph(text, style="List Bullet")
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(10.5)

    later = doc.add_paragraph()
    lrun = later.add_run("Staff Engineer, Contoso  |  Jan 2021 – Present")
    lrun.bold = True
    lrun.font.name = "Calibri"
    lrun.font.size = Pt(11)
    current = doc.add_paragraph(
        "Led platform work on AWS EKS and Terraform for regulated data products.",
        style="List Bullet",
    )
    for run in current.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)

    doc.save(path)


def test_replace_keeps_list_and_font(tmp_path: Path) -> None:
    source = tmp_path / "resume.docx"
    _sample_doc(source)
    doc, rows = parse_document(str(source), present_year=2026)
    bullets = [r for r in rows if r.is_bullet]
    assert len(bullets) == 4
    target = bullets[0]
    apply_bullet_updates(
        doc,
        rows,
        {target.id: "Rebuilt Java 8 Spring Boot intake services and stabilized overnight claims processing."},
    )
    saved = tmp_path / "out.docx"
    doc.save(saved)

    again = Document(str(saved))
    updated = [p for p in again.paragraphs if p.text.startswith("Rebuilt Java 8")]
    assert updated
    paragraph = updated[0]
    assert paragraph.style.name == "List Bullet"
    run = paragraph.runs[0]
    assert run.font.size == Pt(10.5)
    heading = [p for p in again.paragraphs if p.text == "EXPERIENCE"][0]
    assert heading.runs[0].bold is True
    assert heading.runs[0].font.size == Pt(12)
    originals = [
        "Designed PostgreSQL schemas and indexed queries that cut report time for finance operations.",
        "Automated Jenkins pipelines for test, package, and deploy of internal services.",
    ]
    texts = [p.text for p in again.paragraphs]
    for original in originals:
        assert original in texts
    name = [p for p in again.paragraphs if p.text == "Alex Rivera"][0]
    assert name.runs[0].font.color.rgb == RGBColor(0x17, 0x36, 0x5D)
    assert name.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_parse_attaches_role_dates(tmp_path: Path) -> None:
    source = tmp_path / "resume.docx"
    _sample_doc(source)
    _doc, rows = parse_document(str(source), present_year=2026)
    bullets = [r for r in rows if r.is_bullet]
    assert bullets[0].end_year == 2020
    assert "Northwind" in bullets[0].role_title
    assert bullets[-1].is_current_role is True
    assert bullets[-1].end_year == 2026


def test_timeline_blocks_future_tools() -> None:
    from app.tailor import validate_updates
    from app.parse_resume import ResumeParagraph

    row = ResumeParagraph(
        id=3,
        text="Built Java 8 Spring Boot services for claims intake.",
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Staff Engineer, Northwind",
        role_dates="2018–2020",
        start_year=2018,
        end_year=2020,
        is_current_role=False,
        char_budget=52,
    )
    applied, report, warnings = validate_updates(
        [row],
        {
            "updates": [
                {
                    "id": 3,
                    "rewritten": "Built LangGraph agents with GPT-4o on Amazon Bedrock for claims intake.",
                    "changed": True,
                    "reason": "keyword stuffing",
                }
            ]
        },
        strict_timeline=True,
        match_length=False,
    )
    assert applied == {}
    assert report[0]["changed"] is False
    assert warnings and warnings[0]["type"] == "timeline"


def test_length_guard_keeps_layout() -> None:
    from app.tailor import validate_updates
    from app.parse_resume import ResumeParagraph

    original = "Built Java 8 services."
    row = ResumeParagraph(
        id=1,
        text=original,
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Engineer",
        role_dates="2021–Present",
        start_year=2021,
        end_year=2026,
        is_current_role=True,
        char_budget=len(original),
    )
    applied, report, _warnings = validate_updates(
        [row],
        {
            "updates": [
                {
                    "id": 1,
                    "rewritten": original + " " + ("Kubernetes Docker Terraform " * 20),
                    "changed": True,
                    "reason": "too long",
                }
            ]
        },
        strict_timeline=True,
        match_length=True,
    )
    assert applied == {}
    assert report[0]["rewritten"] == original


def test_profile_roundtrip(tmp_path: Path, monkeypatch) -> None:
    import app.profiles as profiles

    monkeypatch.setattr(profiles, "ROOT", tmp_path / "profiles")
    source = tmp_path / "resume.docx"
    _sample_doc(source)
    created = profiles.create_profile("Alex Rivera", str(source), "Alex Rivera.docx")
    assert created["name"] == "Alex Rivera"
    assert created["bullets"] == 4
    listed = profiles.list_profiles()
    assert listed[0]["id"] == created["id"]
    loaded = profiles.get_profile(created["id"])
    assert loaded["filename"] == "Alex Rivera.docx"
    profiles.delete_profile(created["id"])
    assert profiles.list_profiles() == []


def test_comma_dates_and_job_headers_are_not_bullets(tmp_path: Path) -> None:
    from app.dates import parse_role_dates

    dates = parse_role_dates("Lead AI/ML Engineer | AI Architect - IBM  Mar, 2025 – Jul, 2026", 2026)
    assert dates is not None
    assert dates.start_year == 2025
    assert dates.end_year == 2026

    doc = Document()
    doc.add_paragraph("Lead AI/ML Engineer | AI Architect - IBM  Mar, 2025 – Jul, 2026", style="List Bullet")
    doc.add_paragraph(
        "Built Python FastAPI services for model training, validation, and deployment.",
        style="List Bullet",
    )
    path = tmp_path / "headers.docx"
    doc.save(path)
    _doc, rows = parse_document(str(path), present_year=2026)
    bullets = [r for r in rows if r.is_bullet]
    assert len(bullets) == 1
    assert bullets[0].start_year == 2025
    assert bullets[0].end_year == 2026
    assert "IBM" in bullets[0].role_title


def test_extracts_jd_stack_and_weaves_family_sql() -> None:
    from app.keywords import allowed_for_group, extract_jd_stack
    from app.timeline import tech_index

    jd = (
        "We are hiring a full stack AI engineer for a healthcare payer platform. "
        "Build Python FastAPI REST APIs, PostgreSQL, Kubernetes, CI/CD, and prompt engineering. "
        "Required: Python, SQL, Docker. "
        "Preferred qualifications: Salesforce and Azure OpenAI."
    )
    stack = extract_jd_stack(jd)
    assert "Python" in stack["all"]
    assert "FastAPI" in stack["all"]
    assert "REST" in stack["all"]
    assert "SQL" in stack["all"]
    assert "Salesforce" in stack["all"]
    assert "full-stack" in stack["domains"] or "healthcare payer" in stack["domains"]

    allowed = allowed_for_group(
        jd_stack=stack,
        resume_names={"Python", "PostgreSQL", "FastAPI", "Jenkins"},
        role_names={"Python", "PostgreSQL"},
        role_end_year=2026,
        is_skills_or_summary=True,
        tech_by_name=tech_index(),
        is_current_role=True,
    )
    assert "SQL" in allowed["weave"]
    assert "CI/CD" in allowed["weave"]
    assert "Salesforce" in allowed["missing"]


def test_azure_ml_jd_extracts_apis_and_ml_phrases() -> None:
    from app.keywords import extract_jd_stack, inject_missing_keywords
    from app.parse_resume import ResumeParagraph

    jd = (
        "Role: Azure AI/ML Engineer. Must Have Technical Skills: An ML Engineer designs "
        "end to end machine learning solutions across data ingestion, model training, and "
        "production deployment. Develop data pipelines and ML models using Python and "
        "integrate models into application layers (APIs / services). Maintain MLOps / CI CD "
        "pipelines. Good to have: custom co pilots and intelligent agents."
    )
    stack = extract_jd_stack(jd)
    for needed in ("Python", "Azure", "CI/CD", "MLOps", "APIs", "machine learning", "data pipelines"):
        assert needed in stack["all"], needed
    assert "AI Agents" in stack["all"] or "Microsoft Copilot" in stack["all"]

    original = "Built Python FastAPI services for claims intake and model scoring."
    row = ResumeParagraph(
        id=4,
        text=original,
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Technical Skills",
        role_dates="",
        start_year=None,
        end_year=None,
        is_current_role=False,
        char_budget=len(original),
        is_summary=False,
    )
    applied = inject_missing_keywords(
        [row],
        {},
        jd_stack=stack,
        resume_names={"Python", "FastAPI", "Azure"},
        match_length=True,
    )
    blob = applied.get(4, original)
    assert "API" in blob or "MLOps" in blob or "machine learning" in blob.lower() or "CI/CD" in blob


def test_azure_is_not_woven_into_gcp_role() -> None:
    from app.keywords import allowed_for_group
    from app.timeline import tech_index

    allowed = allowed_for_group(
        jd_stack={"required": ["Azure"], "preferred": [], "domains": [], "all": ["Azure"]},
        resume_names={"GCP", "Python", "Azure"},
        role_names={"GCP", "Python"},
        role_end_year=2024,
        is_skills_or_summary=False,
        tech_by_name=tech_index(),
    )
    assert "Azure" not in allowed["weave"]


def test_inject_swaps_database_wording_for_sql() -> None:
    from app.keywords import inject_missing_keywords
    from app.parse_resume import ResumeParagraph

    original = "Built Python services against relational databases for claims intake."
    row = ResumeParagraph(
        id=1,
        text=original,
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Technical Skills",
        role_dates="",
        start_year=None,
        end_year=None,
        is_current_role=False,
        char_budget=len(original),
        is_summary=False,
    )
    applied = inject_missing_keywords(
        [row],
        {},
        jd_stack={"all": ["SQL"], "required": ["SQL"], "preferred": [], "domains": []},
        resume_names={"PostgreSQL"},
        match_length=True,
    )
    assert 1 in applied
    assert "SQL" in applied[1]


def test_llm_payload_canonicalizes_required_stack() -> None:
    from app.keywords import merge_jd_stacks, stack_from_llm_payload

    llm = stack_from_llm_payload(
        {
            "required": ["Python", "CI CD pipelines", "REST APIs", "machine learning", "communication"],
            "preferred": ["Salesforce"],
            "domains": ["healthcare payer"],
        }
    )
    assert "Python" in llm["required"]
    assert "CI/CD" in llm["all"]
    assert "REST" in llm["all"]
    assert "machine learning" in llm["all"]
    assert "Python" in llm["priority"]
    assert "Salesforce" in llm["preferred"]
    assert "communication" not in llm["all"]
    catalog = {
        "required": ["Azure"],
        "preferred": [],
        "domains": ["data pipelines"],
        "all": ["Azure", "data pipelines"],
    }
    merged = merge_jd_stacks(llm, catalog)
    assert "Azure" in merged["all"]
    assert merged["source"] == "llm"


def test_summary_paragraph_is_tailorable(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("PROFESSIONAL SUMMARY")
    doc.add_paragraph(
        "Senior engineer who designs databases and APIs for payer operations across healthcare platforms."
    )
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Engineer, Northwind  |  Jan 2022 – Present")
    doc.add_paragraph("Built Python FastAPI services for intake.", style="List Bullet")
    path = tmp_path / "summary.docx"
    doc.save(path)
    _doc, rows = parse_document(str(path), present_year=2026)
    summaries = [r for r in rows if r.is_summary]
    assert summaries
    assert summaries[0].char_budget > 60
    bullets = [r for r in rows if r.is_bullet]
    assert len(bullets) == 1


def test_parses_fenced_anthropic_json() -> None:
    from app.tailor import _parse_model_json, provider_for

    payload = _parse_model_json(
        """```json
{"target_role":"Engineer","company":"Acme","updates":[{"id":1,"rewritten":"Kept.","reason":"ok"}]}
```"""
    )
    assert payload["updates"][0]["id"] == 1
    assert provider_for("claude-sonnet-4-5", "") == "anthropic"
    assert provider_for("gpt-4o", "anthropic") == "anthropic"
    assert provider_for("gpt-4o", "") == "openai"


def test_openai_quota_is_not_generic_rate_limit() -> None:
    from openai import RateLimitError

    from app.main import provider_http_error

    class _Req:
        def __init__(self) -> None:
            self.method = "POST"
            self.url = "https://api.openai.com/v1/chat/completions"

    class Quota(RateLimitError):
        def __init__(self) -> None:
            Exception.__init__(self, "Error code: 429")
            self.message = "Error code: 429"
            self.body = {
                "error": {
                    "message": "You exceeded your current quota, please check your plan and billing details.",
                    "type": "insufficient_quota",
                    "code": "insufficient_quota",
                }
            }
            self.code = None
            self.status_code = 429
            self.request = _Req()
            self.response = None
            self.request_id = None

    status, detail = provider_http_error(Quota(), "openai")
    assert status == 429
    assert "quota" in detail.lower()
    assert "billing" in detail.lower() or "credits" in detail.lower()


def test_polish_collapses_comma_dumps() -> None:
    from app.keywords import polish_sentence
    from app.parse_resume import ResumeParagraph
    from app.tailor import validate_updates

    assert polish_sentence("Python, Azure,,,, SQL, , CI/CD") == "Python, Azure, SQL, CI/CD"
    original = "Built Python FastAPI services for claims intake."
    row = ResumeParagraph(
        id=1,
        text=original,
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Engineer, Northwind",
        role_dates="2022–Present",
        start_year=2022,
        end_year=2026,
        is_current_role=True,
        char_budget=len(original),
    )
    _applied, report, _warnings = validate_updates(
        [row],
        {
            "updates": [
                {
                    "id": 1,
                    "rewritten": "Python, Azure, SQL,,,,, CI/CD, MLOps, APIs",
                    "reason": "stuffed",
                }
            ]
        },
        strict_timeline=False,
        match_length=False,
    )
    assert ",," not in report[0]["rewritten"]
    assert report[0]["rewritten"] == original


def test_priority_keywords_repeat_across_date_safe_roles() -> None:
    from app.parse_resume import ResumeParagraph
    from app.tailor import validate_updates

    def row(pid: int, text: str, role: str, start: int, end: int, *, skills: bool = False) -> ResumeParagraph:
        return ResumeParagraph(
            id=pid,
            text=text,
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title=role,
            role_dates=f"{start}–{end}",
            start_year=start,
            end_year=end,
            is_current_role=end >= 2026,
            char_budget=len(text),
            is_summary=skills,
        )

    rows = [
        row(1, "Java, Spring Boot, PostgreSQL, Jenkins", "Technical Skills", 2016, 2026, skills=True),
        row(2, "Built Java services for claims intake and reporting.", "Engineer, Northwind", 2016, 2019),
        row(3, "Designed FastAPI services for intake scoring.", "Engineer, Contoso", 2022, 2026),
        row(4, "Automated Jenkins jobs for package and deploy.", "Engineer, Contoso", 2022, 2026),
    ]
    applied, _report, _warnings = validate_updates(
        rows,
        {"updates": []},
        strict_timeline=True,
        match_length=False,
        job_description="Hiring a Python engineer for FastAPI services.",
        jd_stack={
            "required": ["Python"],
            "preferred": [],
            "domains": [],
            "priority": ["Python"],
            "all": ["Python"],
        },
    )
    blob = "\n".join(applied.get(item.id, item.text) for item in rows)
    assert blob.lower().count("python") >= 3
    assert "using Python using" not in blob
    assert ",," not in blob
    northwind = applied.get(2, rows[1].text)
    assert "Python" in northwind


def test_preferred_skills_mentioned_twice_and_required_keeps_existing_count() -> None:
    from app.parse_resume import ResumeParagraph
    from app.tailor import validate_updates

    def row(pid: int, text: str, role: str, start: int, end: int, *, skills: bool = False) -> ResumeParagraph:
        return ResumeParagraph(
            id=pid,
            text=text,
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title=role,
            role_dates=f"{start}–{end}",
            start_year=start,
            end_year=end,
            is_current_role=end >= 2026,
            char_budget=len(text),
            is_summary=skills,
        )

    rows = [
        row(1, "Python, FastAPI, PostgreSQL, Docker", "Technical Skills", 2020, 2026, skills=True),
        row(2, "Built Python FastAPI services with SQL queries for claims.", "Engineer, Contoso", 2022, 2026),
        row(3, "Tuned PostgreSQL and Python jobs for overnight scoring.", "Engineer, Contoso", 2022, 2026),
        row(4, "Packaged Python services for staging deploys.", "Engineer, Northwind", 2018, 2021),
    ]
    applied, _report, _warnings = validate_updates(
        rows,
        {"updates": []},
        strict_timeline=True,
        match_length=False,
        job_description="Required: Python. Preferred: Docker.",
        jd_stack={
            "required": ["Python"],
            "preferred": ["Docker"],
            "domains": [],
            "priority": ["Python"],
            "all": ["Python", "Docker"],
        },
    )
    blob = "\n".join(applied.get(item.id, item.text) for item in rows)
    python_mentions = sum(1 for item in rows if "python" in applied.get(item.id, item.text).lower())
    docker_mentions = sum(1 for item in rows if "docker" in applied.get(item.id, item.text).lower())
    assert python_mentions >= 4
    assert docker_mentions >= 2
    assert ",," not in blob


def test_replace_keeps_body_run_color(tmp_path: Path) -> None:
    from docx.oxml.ns import qn

    doc = Document()
    paragraph = doc.add_paragraph(style="List Bullet")
    lead = paragraph.add_run("Led ")
    lead.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    lead.font.size = Pt(8)
    body = paragraph.add_run("platform work on AWS EKS and Terraform for regulated data products.")
    body.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    body.font.size = Pt(10.5)
    path = tmp_path / "color.docx"
    doc.save(path)

    parsed, rows = parse_document(str(path), present_year=2026)
    target = [row for row in rows if row.is_bullet][0]
    apply_bullet_updates(
        parsed,
        rows,
        {target.id: "Led AWS EKS and Terraform work for regulated data products."},
    )
    saved = tmp_path / "color-out.docx"
    parsed.save(saved)
    again = Document(str(saved))
    updated = [p for p in again.paragraphs if "AWS EKS" in p.text][0]
    run = updated.runs[0]
    assert run.font.size == Pt(10.5)
    assert run.font.color.rgb == RGBColor(0x22, 0x22, 0x22)
    color = run._r.find(qn("w:rPr")).find(qn("w:color"))
    assert color is not None
    assert color.get(qn("w:val")).lower() == "222222"


def test_undated_latest_job_gets_jd_keywords() -> None:
    from app.keywords import inject_missing_keywords
    from app.parse_resume import ResumeParagraph
    from app.tailor import validate_updates

    latest = ResumeParagraph(
        id=1,
        text="Built FastAPI services for claims intake and reporting.",
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Engineer, Contoso",
        role_dates="",
        start_year=None,
        end_year=None,
        is_current_role=False,
        char_budget=54,
    )
    older = ResumeParagraph(
        id=2,
        text="Wrote Java batch jobs for nightly finance files.",
        is_bullet=True,
        style_name="List Bullet",
        in_table=False,
        manual_prefix="",
        role_title="Engineer, Northwind",
        role_dates="2015–2017",
        start_year=2015,
        end_year=2017,
        is_current_role=False,
        char_budget=48,
    )
    applied = inject_missing_keywords(
        [latest, older],
        {},
        jd_stack={
            "required": ["Python"],
            "preferred": [],
            "domains": [],
            "priority": ["Python"],
            "all": ["Python"],
        },
        resume_names={"FastAPI", "Python", "Java"},
        match_length=False,
    )
    assert "Python" in applied.get(1, latest.text)
    assert "Python" not in applied.get(2, older.text)

    _applied, report, _warnings = validate_updates(
        [latest, older],
        {"updates": []},
        strict_timeline=True,
        match_length=False,
        job_description="Hiring a Python engineer.",
        jd_stack={
            "required": ["Python"],
            "preferred": [],
            "domains": [],
            "priority": ["Python"],
            "all": ["Python"],
        },
    )
    latest_text = next(item["rewritten"] for item in report if item["id"] == 1)
    assert "Python" in latest_text


def test_role_profiles_are_configurable_and_hybrid() -> None:
    from app.roles import infer_role_profile, list_role_profiles, resolve_role_profile

    roles = {item["roleId"] for item in list_role_profiles()}
    assert "ai-ml-engineer" in roles
    assert "agentic-ai-engineer" in roles
    assert "software-engineer" in roles
    assert "custom" in roles
    inferred = infer_role_profile("Senior AI/Agentic Engineer for multi-agent LangGraph systems")
    assert inferred["roleId"] == "agentic-ai-engineer"
    hybrid = resolve_role_profile(
        job_description="Principal AI Platform Engineer with Java backend experience and enterprise GenAI architecture.",
        selected_role_id="ai-platform-engineer",
    )
    assert hybrid["roleId"] == "ai-platform-engineer"
    assert "model-serving" in (hybrid.get("writingEmphasis") or "") or "infrastructure" in (hybrid.get("writingEmphasis") or "")


def test_jd_capability_phrases_are_woven_from_related_work() -> None:
    from app.parse_resume import ResumeParagraph
    from app.tailor import validate_updates

    def row(pid: int, text: str, role: str, end: int, *, skills: bool = False) -> ResumeParagraph:
        return ResumeParagraph(
            id=pid,
            text=text,
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title=role,
            role_dates=f"{end-3}–{end}",
            start_year=end - 3,
            end_year=end,
            is_current_role=end >= 2026,
            char_budget=len(text),
            is_summary=skills,
        )

    rows = [
        row(1, "Backend & API Development: Python, FastAPI, Django, REST", "SKILLS", 2026, skills=True),
        row(2, "Implemented multi-agent orchestration in Python using LangChain agents.", "Lead AI/ML Engineer, IBM", 2026),
        row(3, "Built CI/CD around Git with PR reviews and unit/integration tests.", "Lead AI/ML Engineer, IBM", 2026),
        row(4, "Introduced LangGraph to formalize stateful agent execution and checkpointing.", "Senior AI/ML Engineer, PWC", 2025),
    ]
    applied, _report, _warnings = validate_updates(
        rows,
        {"updates": []},
        strict_timeline=True,
        match_length=False,
        job_description="Required: API design, multi-agent coordination, agent architecture, secure SDLC.",
        jd_stack={
            "required": ["API design", "multi-agent coordination", "agent architecture", "secure SDLC"],
            "preferred": ["Temporal"],
            "domains": [],
            "priority": ["API design", "multi-agent coordination"],
            "all": ["API design", "multi-agent coordination", "agent architecture", "secure SDLC", "Temporal"],
        },
    )
    blob = "\n".join(applied.get(item.id, item.text) for item in rows)
    assert "API design" in blob
    assert "multi-agent coordination" in blob
    assert "agent architecture" in blob or "secure SDLC" in blob
    assert "Salesforce" not in blob


def test_extracts_go_from_python_and_go_jd() -> None:
    from app.keywords import canonicalize_keyword, extract_jd_stack, _keyword_in_text
    from app.timeline import technologies_in_text

    jd = (
        "Role: Systems Engineer (MID). Exposure to AWS services. "
        "6–9 years of hands-on experience in Python and Go development. "
        "Experience with Generative AI, MLOps, or AWS SageMaker is an added advantage."
    )
    stack = extract_jd_stack(jd)
    assert "Go" in stack["required"]
    assert "Python" in stack["required"]
    assert canonicalize_keyword("Go") == "Go"
    assert canonicalize_keyword("Golang") == "Go"
    assert any(tech.name == "Go" for tech in technologies_in_text("Python and Go development"))
    assert not any(tech.name == "Go" for tech in technologies_in_text("Google Cloud and ongoing support"))
    assert not _keyword_in_text("Go", "Google Cloud and ongoing Django development")
    assert _keyword_in_text("Go", "Built Python and Go services")
    assert _keyword_in_text("Go", "Production Golang APIs")


def test_weaves_required_go_onto_python_backend_resume() -> None:
    from app.keywords import (
        allowed_for_group,
        ensure_mention_targets,
        extract_jd_stack,
        inject_missing_keywords,
    )
    from app.parse_resume import ResumeParagraph
    from app.timeline import tech_index

    jd = (
        "Systems Engineer. Required: 6-9 years of hands-on experience in Python and Go development. "
        "Preferred: Salesforce, AWS SageMaker."
    )
    stack = extract_jd_stack(jd)
    original = "Built Python/FastAPI services for claims intake and model scoring."
    skills = "Python, FastAPI, AWS, PostgreSQL, Docker"
    rows = [
        ResumeParagraph(
            id=1,
            text=original,
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title="Software Engineer - Microsoft",
            role_dates="2022 – Present",
            start_year=2022,
            end_year=2026,
            is_current_role=True,
            char_budget=len(original),
            is_summary=False,
        ),
        ResumeParagraph(
            id=2,
            text=skills,
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title="Technical Skills",
            role_dates="",
            start_year=None,
            end_year=None,
            is_current_role=False,
            char_budget=len(skills),
            is_summary=False,
        ),
    ]
    allowed = allowed_for_group(
        jd_stack=stack,
        resume_names={"Python", "FastAPI", "AWS", "Java"},
        role_names={"Python", "FastAPI", "AWS"},
        role_end_year=2026,
        is_skills_or_summary=False,
        tech_by_name=tech_index(),
        is_current_role=True,
    )
    assert "Go" in allowed["weave"]
    assert "Salesforce" in allowed["missing"]

    applied = inject_missing_keywords(
        rows,
        {},
        jd_stack=stack,
        resume_names={"Python", "FastAPI", "AWS", "Java"},
        match_length=True,
    )
    applied = ensure_mention_targets(
        rows,
        applied,
        jd_stack=stack,
        resume_names={"Python", "FastAPI", "AWS", "Java"},
        match_length=True,
    )
    blob = "\n".join(applied.get(row.id, row.text) for row in rows)
    assert re.search(r"(?<![A-Za-z0-9])Go(?![A-Za-z0-9])", blob)
    assert "Salesforce" not in blob
    assert "Python/Go/FastAPI" in applied.get(1, original) or "Python and Go" in applied.get(1, original)
    assert "Go" in applied.get(2, skills)


def test_required_go_stays_on_language_and_skills_lines() -> None:
    from app.keywords import ensure_mention_targets, extract_jd_stack, inject_missing_keywords
    from app.parse_resume import ResumeParagraph

    stack = extract_jd_stack(
        "Required: hands-on experience in Python and Go development. Preferred: Salesforce."
    )
    unrelated = "Led stakeholder reviews for claims intake and model scoring."
    rows = [
        ResumeParagraph(
            id=1,
            text="Built Python FastAPI services for claims intake.",
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title="Software Engineer - Microsoft",
            role_dates="2022 – Present",
            start_year=2022,
            end_year=2026,
            is_current_role=True,
            char_budget=len("Built Python FastAPI services for claims intake."),
            is_summary=False,
        ),
        ResumeParagraph(
            id=2,
            text=unrelated,
            is_bullet=True,
            style_name="List Bullet",
            in_table=False,
            manual_prefix="",
            role_title="Software Engineer - Microsoft",
            role_dates="2022 – Present",
            start_year=2022,
            end_year=2026,
            is_current_role=True,
            char_budget=len(unrelated),
            is_summary=False,
        ),
    ]
    applied = inject_missing_keywords(
        rows,
        {},
        jd_stack=stack,
        resume_names={"Python", "FastAPI"},
        match_length=True,
    )
    applied = ensure_mention_targets(
        rows,
        applied,
        jd_stack=stack,
        resume_names={"Python", "FastAPI"},
        match_length=True,
    )
    assert re.search(r"(?<![A-Za-z0-9])Go(?![A-Za-z0-9])", applied.get(1, rows[0].text))
    assert "Go" not in applied.get(2, unrelated)
    assert "Salesforce" not in "\n".join(applied.get(row.id, row.text) for row in rows)



